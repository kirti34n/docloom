"""Regression tests for the studio-ingest audit round:

1. ingest.py:446 -- reingesting a text/research source after its meta['text']
   has already been popped must recover from chunks.jsonl, not fail.
2. sources.py:84 -- an upload whose filename contains a Windows-reserved
   character must be rejected with 400 and leave no orphan source directory.
3. ingest.py:80 -- docx/pptx/xlsx/epub parsers must refuse an archive whose
   declared uncompressed size blows past the configured cap (zip bomb guard),
   while still parsing normal small archives unchanged.
4. ingest.py:156 -- a CSV containing a field larger than csv's default
   131072-char limit must still parse instead of raising csv.Error.
"""

from __future__ import annotations

import asyncio
import io
import json
import os
import tempfile
import zipfile

os.environ.setdefault("DOCLOOM_STUDIO_HOME", tempfile.mkdtemp(prefix="ds-ultracode-ing-"))

import numpy as np  # noqa: E402
import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from docloom_studio import embeddings as E  # noqa: E402
from docloom_studio import ingest as ingest_mod  # noqa: E402
from docloom_studio.db import execute, init_db, new_id, now, query_one  # noqa: E402
from docloom_studio.ingest import _source_dir, load_chunks  # noqa: E402
from docloom_studio.main import app  # noqa: E402
from docloom_studio.settings import data_dir  # noqa: E402


@pytest.fixture(autouse=True)
def _db():
    init_db()


# ================================================== 1. text/research reingest


def _notebook() -> str:
    nb = new_id()
    execute("INSERT INTO notebooks (id, name, created, updated) VALUES (?, ?, ?, ?)",
            (nb, "nb", now(), now()))
    return nb


def _text_source(nb: str, title: str, text: str) -> str:
    sid = new_id()
    execute("INSERT INTO sources (id, notebook_id, kind, title, status, "
            "context_mode, meta_json, created) VALUES (?, ?, 'text', ?, 'pending', "
            "'full', ?, ?)", (sid, nb, title, json.dumps({"text": text}), now()))
    return sid


def test_reingest_text_source_recovers_from_chunks_after_meta_text_popped(monkeypatch):
    async def fake_embed_source(source_id, texts, name="embeddings"):
        np.save(_source_dir(source_id) / f"{name}.npy",
                np.zeros((len(texts), 8), dtype=np.float32))

    async def fake_summarize(source_id, chunks):
        return ""

    monkeypatch.setattr(E, "embed_source", fake_embed_source)
    monkeypatch.setattr(ingest_mod, "_summarize_source", fake_summarize)

    nb = _notebook()
    body = ("Blue whales are the largest animals ever known to have lived, "
            "reaching 30 metres in length. They feed almost exclusively on krill.")
    sid = _text_source(nb, "Whales", body)

    asyncio.run(ingest_mod.ingest_source(sid))
    row = query_one("SELECT status, meta_json FROM sources WHERE id = ?", (sid,))
    assert row["status"] == "ready"
    assert "text" not in json.loads(row["meta_json"])  # popped as designed
    first_chunks = load_chunks(sid)
    assert first_chunks

    # The documented recovery path for a 'stale' source: flip back to pending
    # and re-ingest. meta['text'] is already gone -- before the fix this hit
    # `if not chunks: raise ValueError("no extractable text")` and the source
    # was permanently destroyed instead of recovered.
    execute("UPDATE sources SET status = 'pending' WHERE id = ?", (sid,))
    asyncio.run(ingest_mod.ingest_source(sid))
    row2 = query_one("SELECT status, meta_json FROM sources WHERE id = ?", (sid,))
    assert row2["status"] == "ready"
    assert json.loads(row2["meta_json"]).get("error") is None
    assert load_chunks(sid) == first_chunks


def test_reingest_research_source_also_recovers_from_chunks(monkeypatch):
    async def fake_embed_source(source_id, texts, name="embeddings"):
        np.save(_source_dir(source_id) / f"{name}.npy",
                np.zeros((len(texts), 8), dtype=np.float32))

    async def fake_summarize(source_id, chunks):
        return ""

    monkeypatch.setattr(E, "embed_source", fake_embed_source)
    monkeypatch.setattr(ingest_mod, "_summarize_source", fake_summarize)

    nb = _notebook()
    sid = new_id()
    execute("INSERT INTO sources (id, notebook_id, kind, title, status, "
            "context_mode, meta_json, created) VALUES (?, ?, 'research', ?, 'pending', "
            "'full', ?, ?)",
            (sid, nb, "Research", json.dumps({"text": "Findings about volcano types. " * 5}), now()))

    asyncio.run(ingest_mod.ingest_source(sid))
    assert query_one("SELECT status FROM sources WHERE id = ?", (sid,))["status"] == "ready"

    execute("UPDATE sources SET status = 'pending' WHERE id = ?", (sid,))
    asyncio.run(ingest_mod.ingest_source(sid))
    row = query_one("SELECT status FROM sources WHERE id = ?", (sid,))
    assert row["status"] == "ready"
    assert load_chunks(sid)


# ======================================== 2. Windows-reserved-char filenames


def _register_notebook() -> tuple:
    client = TestClient(app)
    client.__enter__()
    client.post("/api/auth/register",
                json={"email": f"{new_id()}@t.local", "password": "password123"})
    ws = client.get("/api/workspaces").json()
    nb = client.post("/api/notebooks",
                     json={"name": "n", "workspace_id": ws[0]["id"]}).json()["id"]
    return client, nb


def test_upload_rejects_reserved_chars_with_no_orphan_dir_then_allows_normal_name():
    client, nb = _register_notebook()
    try:
        sources_root = data_dir() / "sources"
        before = {p.name for p in sources_root.iterdir()}

        for bad_name in ("report|2024.txt", "notes<draft>.txt", "glob*.md"):
            r = client.post(f"/api/notebooks/{nb}/sources/file",
                            files={"file": (bad_name, io.BytesIO(b"hello"), "text/plain")})
            assert r.status_code == 400, (bad_name, r.status_code, r.text)

        after = {p.name for p in sources_root.iterdir()}
        assert after == before  # no orphan source directories left behind

        ok = client.post(f"/api/notebooks/{nb}/sources/file",
                         files={"file": ("ok.txt", io.BytesIO(b"hello"), "text/plain")})
        assert ok.status_code == 200, ok.text
        sid = ok.json()["source_id"]
        assert (_source_dir(sid) / "ok.txt").is_file()
    finally:
        client.__exit__(None, None, None)


# ==================================================== 3. zip-bomb parser cap


def test_parse_docx_rejects_declared_oversize_archive(tmp_path, monkeypatch):
    bomb = tmp_path / "bomb.docx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", "A" * (2 * 1024 * 1024))

    monkeypatch.setattr(ingest_mod, "MAX_ZIP_UNCOMPRESSED_BYTES", 1024 * 1024)
    with pytest.raises(ValueError):
        ingest_mod.parse_docx(bomb)


def test_parse_docx_still_parses_a_normal_small_docx(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Ship the podcast feature")
    p = tmp_path / "real.docx"
    doc.save(str(p))
    text = ingest_mod.parse_docx(p)
    assert "Ship the podcast feature" in text


# ================================================ 4. oversized single CSV field


def test_parse_csv_field_over_default_limit_does_not_raise():
    huge = "x" * 200000
    out = ingest_mod.parse_csv(f"a,{huge}")
    assert huge in out


# ============================ 5. URL ingest: body cap + DNS-rebind peer re-check


def test_fetch_url_caps_an_oversized_body(monkeypatch):
    """A public host that streams more than _MAX_FETCH_BYTES must be rejected
    instead of buffered whole into memory."""
    import http.server
    import threading

    oversize = ingest_mod._MAX_FETCH_BYTES + 1024

    class _Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            self.send_response(200)
            self.send_header("Content-Length", str(oversize))
            self.end_headers()
            sent = 0
            block = b"x" * (1024 * 1024)
            while sent < oversize:
                self.wfile.write(block)
                sent += len(block)

        def log_message(self, *a):
            pass

    srv = http.server.HTTPServer(("127.0.0.1", 0), _Handler)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    port = srv.server_address[1]
    # bypass the SSRF guard so we exercise the body cap, not the loopback reject
    monkeypatch.setattr(ingest_mod, "_guard_url", lambda url: None)
    monkeypatch.setattr(ingest_mod, "_reject_non_public_peer", lambda resp: None)
    try:
        with pytest.raises(ValueError, match="MB limit"):
            ingest_mod.fetch_url(f"http://127.0.0.1:{port}/big")
    finally:
        srv.shutdown()


def test_reject_non_public_peer_catches_a_rebind_to_loopback():
    """Even when the pre-connect guard passed, a connection whose real peer
    address is internal (a DNS rebind) must be refused before the body is read."""
    class _Stream:
        def get_extra_info(self, key):
            return ("127.0.0.1", 80) if key == "server_addr" else None

    class _Resp:
        extensions = {"network_stream": _Stream()}

    with pytest.raises(ValueError, match="non-public address"):
        ingest_mod._reject_non_public_peer(_Resp())

    class _PublicStream:
        def get_extra_info(self, key):
            return ("8.8.8.8", 80) if key == "server_addr" else None

    class _PublicResp:
        extensions = {"network_stream": _PublicStream()}

    ingest_mod._reject_non_public_peer(_PublicResp())  # must not raise
